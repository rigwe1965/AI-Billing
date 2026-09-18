import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from pypdf import PdfReader

from app.models.budget import Budget
from app.models.pricing_rate import PricingRate
from app.models.usage_record import UsageRecord, UsageSourceType
from app.services.costing_engine import price_usage_records
from app.services.report_service import build_excel_report, build_pdf_report


def _seed_priced_usage(db, provider, department):
    db.add(
        PricingRate(
            provider_id=provider.id,
            sku_code="claude-sonnet-input-tokens",
            unit="per_1m_tokens",
            rate_amount=Decimal("3.00"),
            effective_from=date(2026, 1, 1),
            source="manual_admin_entry",
        )
    )
    db.commit()

    usage = UsageRecord(
        provider_id=provider.id,
        department_id=department.id,
        sku_code="claude-sonnet-input-tokens",
        quantity=Decimal("2000000"),
        unit="tokens",
        period_start=date(2026, 1, 10),
        period_end=date(2026, 1, 10),
        source_type=UsageSourceType.csv_import,
        source_reference="test",
    )
    db.add(usage)
    db.commit()
    db.refresh(usage)
    price_usage_records(db, [usage])


def test_excel_report_has_expected_columns_and_sums(db_session, anthropic_provider, engineering_dept):
    _seed_priced_usage(db_session, anthropic_provider, engineering_dept)

    content = build_excel_report(db_session)
    wb = load_workbook(io.BytesIO(content))

    assert "Summary" in wb.sheetnames
    assert "Detail" in wb.sheetnames

    summary_ws = wb["Summary"]
    header = [c.value for c in summary_ws[1]]
    assert header == ["Period", "Department", "Provider", "Total Amount", "Currency"]
    data_row = [c.value for c in summary_ws[2]]
    assert data_row[1] == "Engineering"
    assert data_row[2] == "Anthropic"
    assert data_row[3] == 6.0  # 2,000,000 tokens * $3/1M

    detail_ws = wb["Detail"]
    detail_header = [c.value for c in detail_ws[1]]
    assert "Pricing Rate ID" in detail_header
    assert "Source Type" in detail_header


def test_excel_report_scoped_to_one_department(db_session, anthropic_provider, engineering_dept):
    from app.models.department import Department

    marketing = Department(name="Marketing", code="MKT")
    db_session.add(marketing)
    db_session.commit()

    _seed_priced_usage(db_session, anthropic_provider, engineering_dept)
    _seed_priced_usage(db_session, anthropic_provider, marketing)

    content = build_excel_report(db_session, department_id=engineering_dept.id)
    wb = load_workbook(io.BytesIO(content))
    summary_ws = wb["Summary"]

    departments_seen = {row[1].value for row in summary_ws.iter_rows(min_row=2)}
    assert departments_seen == {"Engineering"}


def test_pdf_report_renders_without_error_and_has_pages(db_session, anthropic_provider, engineering_dept):
    _seed_priced_usage(db_session, anthropic_provider, engineering_dept)
    db_session.add(Budget(department_id=engineering_dept.id, period="2026-01", amount=Decimal("100")))
    db_session.commit()

    content = build_pdf_report(db_session)

    assert content[:4] == b"%PDF"
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "AI Spend Summary Report" in text


def test_pdf_report_with_no_budgets_does_not_crash(db_session, anthropic_provider, engineering_dept):
    _seed_priced_usage(db_session, anthropic_provider, engineering_dept)
    content = build_pdf_report(db_session)
    assert content[:4] == b"%PDF"
